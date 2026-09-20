import io
import json

import pytest
from fastapi.testclient import TestClient

from app.ai.client import LLMResponse, StubLLMClient
from app.core.exceptions import ResumeExtractionError, ValidationError
from app.db.database import SessionLocal
from app.services.resume_service import extract_resume_text

RESUME_TEXT = """
Alex Rivera
Backend Engineer
Summary: Backend engineer with 4 years building data products.
Skills: Python, SQL, React, Machine Learning
Technologies: FastAPI, PostgreSQL, scikit-learn, Docker
Projects:
- ML Sales Forecaster: built a time-series model forecasting retail sales, deployed with FastAPI.
- Inventory Dashboard: React dashboard over a SQL warehouse.
Experience:
- Backend Engineer, Acme Corp, 2021-2025: owned the data pipeline and API layer.
Education:
- BSc Computer Science, State University, 2021
Certifications: AWS Certified Developer
"""

RESUME_STRUCTURED = json.dumps(
    {
        "name": "Alex Rivera",
        "title": "Backend Engineer",
        "summary": "Backend engineer with 4 years building data products.",
        "skills": ["Python", "SQL", "React", "Machine Learning"],
        "technologies": ["FastAPI", "PostgreSQL", "scikit-learn", "Docker"],
        "projects": [
            {
                "name": "ML Sales Forecaster",
                "description": "Built a time-series model forecasting retail sales, deployed with FastAPI.",
                "technologies": ["Python", "scikit-learn", "FastAPI"],
            },
            {
                "name": "Inventory Dashboard",
                "description": "React dashboard over a SQL warehouse.",
                "technologies": ["React", "SQL"],
            },
        ],
        "experience": [
            {
                "role": "Backend Engineer",
                "company": "Acme Corp",
                "duration": "2021-2025",
                "summary": "Owned the data pipeline and API layer.",
            }
        ],
        "education": [
            {"degree": "BSc Computer Science", "institution": "State University", "year": "2021"}
        ],
        "certifications": ["AWS Certified Developer"],
    }
)

ROLE_QUESTION = json.dumps(
    {
        "question": "Walk me through the ML project on your resume.",
        "question_type": "technical",
        "display_hint": "",
        "expected_duration_seconds": 120,
        "notes": "",
    }
)

REPLACEMENT_QUESTION = json.dumps(
    {
        "question": "How did you decide between forecasting approaches for the model?",
        "question_type": "probe",
        "display_hint": "",
        "expected_duration_seconds": 120,
        "notes": "",
    }
)

FOLLOW_UP_QUESTION = json.dumps(
    {
        "question": "Which deployment tradeoffs did you consider for the model?",
        "question_type": "probe",
        "display_hint": "",
        "expected_duration_seconds": 120,
        "notes": "",
        "performance": "strong",
        "next_difficulty": "hard",
    }
)

ROLE_EVALUATION = json.dumps(
    {
        "overall_score": 81,
        "skills": [
            {"skill": "Technical Knowledge", "score": 85, "evidence": "Explained the model end to end"},
            {"skill": "Communication", "score": 80, "evidence": "Clear explanation"},
            {"skill": "Problem Solving", "score": 78, "evidence": "Traded off complexity vs accuracy"},
            {"skill": "Confidence", "score": 82, "evidence": "Steady and concrete"},
            {"skill": "Resume Alignment", "score": 88, "evidence": "Answers matched resume projects"},
            {"skill": "Completeness", "score": 75, "evidence": "Missed deployment details at first"},
        ],
        "strengths": ["Strong technical depth", "Grounded answers in real projects"],
        "improvements": ["Give more structured walkthroughs"],
        "recommendations": ["Practice system design interviews"],
    }
)


class FakeLLM:
    def __init__(self, responses):
        self._responses = list(responses)
        self.calls = 0
        self.prompts: list[str] = []

    def chat(self, **kwargs):
        self.calls += 1
        self.prompts.append(str(kwargs.get("user_prompt") or ""))
        content = self._responses.pop(0) if self._responses else "{}"
        return LLMResponse(content=content, model="fake-model", tokens_used=10)


def _patch_interview(monkeypatch, fake):
    monkeypatch.setattr("app.services.interview_session_service.get_llm_client", lambda: fake)
    monkeypatch.setattr("app.ai.interviewer.get_llm_client", lambda: fake)


def _patch_evaluator(monkeypatch, fake):
    monkeypatch.setattr("app.services.case_evaluation_service.get_llm_client", lambda: fake)


def _patch_parser(monkeypatch, fake):
    monkeypatch.setattr("app.ai.resume_parser.get_llm_client", lambda: fake)


# ---------------------------------------------------------------------------
# Resume text extraction
# ---------------------------------------------------------------------------


def test_extract_docx_text():
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Alex Rivera")
    doc.add_paragraph("Backend Engineer with SQL experience.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "FastAPI"
    doc.save(buf)

    text = extract_resume_text("resume.docx", buf.getvalue())
    assert "Alex Rivera" in text
    assert "SQL" in text
    assert "Python" in text


def test_extract_docx_corrupt_file():
    """Bytes that merely claim to be a .docx must be rejected cleanly.

    Content sniffing now rejects this *before* python-docx is invoked (a ZIP
    header is required), so the error is a ValidationError rather than a parser
    ResumeExtractionError. Both are 422s to the caller — the point of the test
    is that nothing is fabricated and no parser exception escapes.
    """
    with pytest.raises((ValidationError, ResumeExtractionError)):
        extract_resume_text("resume.docx", b"this is not a real docx")


def test_pdf_extraction_errors_cleanly_when_pypdf_missing():
    """Without pypdf installed the service must fail clearly, not fabricate."""
    try:
        import pypdf  # noqa: F401

        pytest.skip("pypdf installed — real PDF parsing is covered separately")
    except ImportError:
        with pytest.raises(ResumeExtractionError, match="pypdf"):
            extract_resume_text("resume.pdf", b"%PDF-1.4 fake")


def test_pdf_extraction_path_with_fake_module(monkeypatch):
    """Exercise the pypdf code path with an injected fake module so the
    extraction logic is verified even when pypdf is not installed locally."""
    import sys
    import types

    class FakePage:
        def __init__(self, text):
            self._text = text

        def extract_text(self):
            return self._text

    class FakePdfReader:
        def __init__(self, stream):
            raw = stream.read()
            if b"PDF" not in raw:
                raise Exception("bad pdf")
            self.pages = [FakePage("Alex Rivera\nSQL, Python, Tableau")]

    fake = types.ModuleType("pypdf")
    fake.PdfReader = FakePdfReader
    monkeypatch.setitem(sys.modules, "pypdf", fake)

    text = extract_resume_text("resume.pdf", b"%PDF-1.4 valid-ish bytes")
    assert "Alex Rivera" in text
    assert "SQL" in text

    # Content sniffing is the outer layer: a .pdf whose bytes are not a PDF at
    # all never reaches the parser. Defence in depth — see test_security.py for
    # the payload-level cases.
    with pytest.raises(ValidationError):
        extract_resume_text("resume.pdf", b"not a pdf at all")


def test_unsupported_extension_rejected():
    with pytest.raises(Exception) as exc_info:
        extract_resume_text("resume.exe", b"not a resume")
    assert exc_info.value.status_code == 422


# ---------------------------------------------------------------------------
# Resume upload / parse API
# ---------------------------------------------------------------------------


def test_resume_parse_requires_provider_key(client: TestClient, auth_headers, monkeypatch):
    _patch_parser(monkeypatch, StubLLMClient())
    resp = client.post(
        "/api/ai/resume/parse",
        headers=auth_headers,
        files={"file": ("resume.txt", RESUME_TEXT.encode("utf-8"), "text/plain")},
        data={"role": "Software Engineer"},
    )
    assert resp.status_code == 502
    assert "GROQ_API_KEY" in resp.json()["detail"]


def test_resume_parse_returns_structured_data(
    client: TestClient, auth_headers, monkeypatch
):
    _patch_parser(monkeypatch, FakeLLM([RESUME_STRUCTURED]))
    resp = client.post(
        "/api/ai/resume/parse",
        headers=auth_headers,
        files={"file": ("resume.txt", RESUME_TEXT.encode("utf-8"), "text/plain")},
        data={"role": "Software Engineer"},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["source_type"] == "txt"
    assert body["resume"]["name"] == "Alex Rivera"
    assert "Machine Learning" in body["resume"]["skills"]
    projects = body["resume"]["projects"]
    assert any(p["name"] == "ML Sales Forecaster" for p in projects)
    assert body["resume"]["experience"][0]["company"] == "Acme Corp"


def test_resume_parse_rejects_unknown_extension(client: TestClient, auth_headers):
    resp = client.post(
        "/api/ai/resume/parse",
        headers=auth_headers,
        files={"file": ("resume.exe", b"data", "application/octet-stream")},
    )
    assert resp.status_code == 422
    assert "Unsupported" in resp.json()["detail"]


def test_resume_parse_requires_auth(client: TestClient):
    resp = client.post(
        "/api/ai/resume/parse",
        files={"file": ("resume.txt", RESUME_TEXT.encode("utf-8"), "text/plain")},
    )
    assert resp.status_code == 401


# ---------------------------------------------------------------------------
# Role/resume interview chat
# ---------------------------------------------------------------------------


def _start_role_interview(client: TestClient, auth_headers, monkeypatch, responses=None):
    fake = FakeLLM(responses or [ROLE_QUESTION, FOLLOW_UP_QUESTION])
    _patch_interview(monkeypatch, fake)
    resp = client.post(
        "/api/ai/interview/chat",
        headers=auth_headers,
        json={
            "role": "Software Engineer",
            "topic": "Software Engineer",
            "difficulty": "Medium",
            "resume": json.loads(RESUME_STRUCTURED),
            "message": "",
        },
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    session = client.get(f"/api/ai/sessions/{body['session_id']}", headers=auth_headers).json()
    assert session["metadata_"]["mode"] == "role"
    assert session["metadata_"]["role"] == "Software Engineer"
    assert session["metadata_"]["resume"]["projects"][0]["name"] == "ML Sales Forecaster"
    return fake, body["session_id"]


def test_role_interview_session_created_with_resume(
    client: TestClient, auth_headers, monkeypatch
):
    fake, sid = _start_role_interview(client, auth_headers, monkeypatch)
    msgs = client.get(f"/api/ai/sessions/{sid}/messages", headers=auth_headers).json()
    assert [m["role"] for m in msgs] == ["interviewer"]

    # The interviewer prompt must contain the resume (skills + projects) so it
    # can ask about the candidate's actual work — not generic questions.
    assert "Machine Learning" in fake.prompts[0]
    assert "ML Sales Forecaster" in fake.prompts[0]
    assert "Backend Engineer" in fake.prompts[0]
    assert "Acme Corp" in fake.prompts[0]


def test_role_interview_adapts_with_history(client: TestClient, auth_headers, monkeypatch):
    fake, sid = _start_role_interview(
        client, auth_headers, monkeypatch, [ROLE_QUESTION, FOLLOW_UP_QUESTION]
    )
    answer = "I built a time-series model in scikit-learn and served it through FastAPI."
    resp = client.post(
        "/api/ai/interview/chat",
        headers=auth_headers,
        json={"session_id": sid, "message": answer},
    )
    assert resp.status_code == 200, resp.text
    msgs = client.get(f"/api/ai/sessions/{sid}/messages", headers=auth_headers).json()
    roles = [m["role"] for m in msgs]
    assert roles == ["interviewer", "candidate", "interviewer"]
    assert msgs[1]["content"] == answer
    # Follow-up prompt sees both the resume and the candidate's latest answer.
    assert answer in fake.prompts[1]
    assert "ML Sales Forecaster" in fake.prompts[1]


def test_interview_exposes_question_progress_from_config(
    client: TestClient, auth_headers, monkeypatch
):
    fake, sid = _start_role_interview(client, auth_headers, monkeypatch)
    # Start call reports question 1 of the configured total (6 default for
    # open-ended role interviews) and the session persists that target.
    assert fake.calls == 1
    session = client.get(f"/api/ai/sessions/{sid}", headers=auth_headers).json()
    assert session["metadata_"]["total_questions"] == 6
    assert session["metadata_"]["difficulty"] == "Medium"


def test_adaptivity_performance_signal_persisted(
    client: TestClient, auth_headers, monkeypatch
):
    fake, sid = _start_role_interview(
        client, auth_headers, monkeypatch, [ROLE_QUESTION, FOLLOW_UP_QUESTION]
    )
    resp = client.post(
        "/api/ai/interview/chat",
        headers=auth_headers,
        json={"session_id": sid, "message": "I deployed the model behind a REST API."},
    )
    assert resp.status_code == 200, resp.text
    session = client.get(f"/api/ai/sessions/{sid}", headers=auth_headers).json()
    signals = session["metadata_"]["performance_history"]
    assert signals[-1]["performance"] == "strong"
    assert signals[-1]["next_difficulty"] == "Hard"
    # Server-authoritative difficulty adaptation for the next turn.
    assert session["metadata_"]["difficulty"] == "Hard"


def test_duplicate_interview_questions_are_regenerated(
    client: TestClient, auth_headers, monkeypatch
):
    # Second turn first returns the SAME question as the opening question
    # (duplicate) and then a genuinely different one after the bounded retry.
    fake, sid = _start_role_interview(
        client,
        auth_headers,
        monkeypatch,
        [ROLE_QUESTION, ROLE_QUESTION, REPLACEMENT_QUESTION],
    )
    resp = client.post(
        "/api/ai/interview/chat",
        headers=auth_headers,
        json={"session_id": sid, "message": "I trained the model on three years of sales data."},
    )
    assert resp.status_code == 200, resp.text
    assert fake.calls == 3  # opening + duplicate attempt + replacement
    msgs = client.get(f"/api/ai/sessions/{sid}/messages", headers=auth_headers).json()
    questions = [m["content"] for m in msgs if m["role"] == "interviewer"]
    assert len(questions) == 2
    assert questions[0] != questions[1]
    assert "forecasting approaches" in questions[1]


def test_saved_resume_asset_can_start_interview(
    client: TestClient, auth_headers, monkeypatch
):
    # Save a resume asset first, then reference it by id when starting.
    save = client.post(
        "/api/resumes",
        headers=auth_headers,
        json={
            "filename": "alex_resume.txt",
            "role": "Software Engineer",
            "resume": json.loads(RESUME_STRUCTURED),
        },
    )
    assert save.status_code == 201, save.text
    resume_id = save.json()["id"]

    fake = FakeLLM([ROLE_QUESTION])
    _patch_interview(monkeypatch, fake)
    resp = client.post(
        "/api/ai/interview/chat",
        headers=auth_headers,
        json={
            "role": "Software Engineer",
            "topic": "Software Engineer",
            "resume_id": resume_id,
            "message": "",
        },
    )
    assert resp.status_code == 200, resp.text
    session = client.get(f"/api/ai/sessions/{resp.json()['session_id']}", headers=auth_headers).json()
    md = session["metadata_"]
    assert md["resume_id"] == resume_id
    assert md["resume"]["projects"][0]["name"] == "ML Sales Forecaster"
    assert "ML Sales Forecaster" in fake.prompts[0]


def test_other_users_resume_asset_cannot_start_interview(
    client: TestClient, auth_headers, monkeypatch
):
    save = client.post(
        "/api/resumes",
        headers=auth_headers,
        json={
            "filename": "alex_resume.txt",
            "resume": json.loads(RESUME_STRUCTURED),
        },
    ).json()
    other = client.post(
        "/api/auth/signup",
        json={"full_name": "Other", "email": "resume-use-other@example.com", "password": "Password123!", "terms_accepted": True, "privacy_accepted": True},
    ).json()
    other_headers = {"Authorization": f"Bearer {other['access_token']}"}

    _patch_interview(monkeypatch, FakeLLM([ROLE_QUESTION]))
    resp = client.post(
        "/api/ai/interview/chat",
        headers=other_headers,
        json={"role": "Software Engineer", "resume_id": save["id"], "message": ""},
    )
    assert resp.status_code == 404


def test_feedback_and_recommendations_persisted_for_review(
    client: TestClient, auth_headers, monkeypatch
):
    from app.ai.client import StubLLMClient

    interviewer_fake = FakeLLM([ROLE_QUESTION, FOLLOW_UP_QUESTION])
    _patch_interview(monkeypatch, interviewer_fake)
    first = client.post(
        "/api/ai/interview/chat",
        headers=auth_headers,
        json={
            "role": "Software Engineer",
            "topic": "Software Engineer",
            "resume": json.loads(RESUME_STRUCTURED),
            "message": "",
        },
    ).json()
    sid = first["session_id"]
    client.post(
        "/api/ai/interview/chat",
        headers=auth_headers,
        json={"session_id": sid, "message": "A solid answer about the model."},
    )

    _patch_evaluator(monkeypatch, FakeLLM([ROLE_EVALUATION]))
    assert (
        client.post("/api/ai/evaluation", headers=auth_headers, json={"session_id": sid}).status_code
        == 200
    )
    # Deterministic feedback path (no provider) + LLM-free recommendations.
    monkeypatch.setattr(
        "app.ai.feedback_generator.get_llm_client", lambda: StubLLMClient()
    )
    fb = client.post("/api/ai/feedback", headers=auth_headers, json={"session_id": sid})
    assert fb.status_code == 200, fb.text
    recs = client.post(
        "/api/ai/recommendations", headers=auth_headers, json={"session_id": sid}
    )
    assert recs.status_code == 200, recs.text

    # Both are persisted on the session, so reviewing history costs no AI calls.
    session = client.get(f"/api/ai/sessions/{sid}", headers=auth_headers).json()
    md = session["metadata_"]
    assert md["feedback"]["overall_score"] == 81
    assert md["feedback"]["biggest_opportunity"]
    assert "next_best_action" in md["recommendations"]


def test_role_interview_ownership(client: TestClient, auth_headers, monkeypatch):
    _fake, sid = _start_role_interview(client, auth_headers, monkeypatch)
    other_resp = client.post(
        "/api/auth/signup",
        json={"full_name": "Other", "email": "role-other@example.com", "password": "Password123!", "terms_accepted": True, "privacy_accepted": True},
    )
    other_headers = {"Authorization": f"Bearer {other_resp.json()['access_token']}"}

    assert (
        client.get(f"/api/ai/sessions/{sid}", headers=other_headers).status_code == 403
    )
    assert (
        client.get(f"/api/ai/sessions/{sid}/messages", headers=other_headers).status_code == 403
    )
    # Cannot submit answers into another user's session either.
    resp = client.post(
        "/api/ai/interview/chat",
        headers=other_headers,
        json={"session_id": sid, "message": "intruder answer"},
    )
    assert resp.status_code == 403


def test_role_interview_chat_requires_auth(client: TestClient, monkeypatch):
    _patch_interview(monkeypatch, FakeLLM([ROLE_QUESTION]))
    resp = client.post(
        "/api/ai/interview/chat",
        json={"role": "Software Engineer", "message": ""},
    )
    assert resp.status_code == 401


# ---------------------------------------------------------------------------
# Role/resume evaluation → feedback → recommendations → progress
# ---------------------------------------------------------------------------


def test_role_interview_evaluation_flow(
    client: TestClient, auth_headers, monkeypatch
):
    from app.models.profile import Profile
    from app.models.skill import Skill, UserSkill
    from app.models.user import User

    interviewer_fake = FakeLLM([ROLE_QUESTION, FOLLOW_UP_QUESTION])
    _patch_interview(monkeypatch, interviewer_fake)
    first = client.post(
        "/api/ai/interview/chat",
        headers=auth_headers,
        json={
            "role": "Software Engineer",
            "topic": "Software Engineer",
            "resume": json.loads(RESUME_STRUCTURED),
            "message": "",
        },
    )
    sid = first.json()["session_id"]
    second = client.post(
        "/api/ai/interview/chat",
        headers=auth_headers,
        json={"session_id": sid, "message": "I built a forecast model with scikit-learn."},
    )
    assert second.status_code == 200, second.text

    _patch_evaluator(monkeypatch, FakeLLM([ROLE_EVALUATION]))
    eval_resp = client.post(
        "/api/ai/evaluation", headers=auth_headers, json={"session_id": sid}
    )
    assert eval_resp.status_code == 200, eval_resp.text
    evaluation = eval_resp.json()["evaluation"]
    assert evaluation["overall_score"] == 81
    skill_names = {s["skill"] for s in evaluation["skills"]}
    assert "Technical Knowledge" in skill_names
    assert "Resume Alignment" in skill_names

    # Persisted + session marked completed with the evaluation stored.
    session = client.get(f"/api/ai/sessions/{sid}", headers=auth_headers).json()
    assert session["status"] == "completed"
    assert session["metadata_"]["evaluation"]["overall_score"] == 81

    # Progress: real interview score metrics + recorded skills.
    progress = client.get("/api/progress", headers=auth_headers).json()
    assert progress["total_interviews_completed"] == 1
    assert progress["average_interview_score"] == 81.0

    with SessionLocal() as s:
        user = s.query(User).filter(User.email == "test-user-fixture@example.com").first()
        joined = (
            s.query(UserSkill, Skill)
            .join(Skill, UserSkill.skill_id == Skill.id)
            .filter(UserSkill.user_id == user.id)
            .all()
        )
        names = {skill.name for _, skill in joined}
        assert "Technical Knowledge" in names
        assert "Resume Alignment" in names
        profile = s.query(Profile).filter(Profile.user_id == user.id).first()
        assert profile.readiness_score > 0

    # Feedback (deterministic path when the provider is not configured).
    monkeypatch.setattr(
        "app.ai.feedback_generator.get_llm_client", lambda: StubLLMClient()
    )
    fb = client.post("/api/ai/feedback", headers=auth_headers, json={"session_id": sid})
    assert fb.status_code == 200, fb.text
    assert fb.json()["overall_score"] == 81

    # Recommendations based on the real evaluation.
    recs = client.post(
        "/api/ai/recommendations", headers=auth_headers, json={"session_id": sid}
    )
    assert recs.status_code == 200, recs.text
    assert "next_best_action" in recs.json()


def test_role_interview_evaluation_ownership(client: TestClient, auth_headers, monkeypatch):
    interviewer_fake = FakeLLM([ROLE_QUESTION])
    _patch_interview(monkeypatch, interviewer_fake)
    first = client.post(
        "/api/ai/interview/chat",
        headers=auth_headers,
        json={"role": "Data Analyst", "message": ""},
    )
    sid = first.json()["session_id"]

    other_resp = client.post(
        "/api/auth/signup",
        json={"full_name": "Other", "email": "role-eval-other@example.com", "password": "Password123!", "terms_accepted": True, "privacy_accepted": True},
    )
    other_headers = {"Authorization": f"Bearer {other_resp.json()['access_token']}"}
    resp = client.post(
        "/api/ai/evaluation", headers=other_headers, json={"session_id": sid}
    )
    assert resp.status_code == 403


def test_role_interview_without_resume_is_generic_role(
    client: TestClient, auth_headers, monkeypatch
):
    fake = FakeLLM([ROLE_QUESTION])
    _patch_interview(monkeypatch, fake)
    resp = client.post(
        "/api/ai/interview/chat",
        headers=auth_headers,
        json={"role": "Product Manager", "topic": "Product Manager", "message": ""},
    )
    assert resp.status_code == 200
    session = client.get(f"/api/ai/sessions/{resp.json()['session_id']}", headers=auth_headers).json()
    assert session["metadata_"]["mode"] == "role"
    assert "Product Manager" in fake.prompts[0]
    assert "JOB ROLE" in fake.prompts[0]
